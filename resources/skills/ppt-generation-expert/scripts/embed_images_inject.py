#!/usr/bin/env python3
"""Inject human-reviewed evidence images from a controlled folder into a merged pptx.

Controlled-image contract (Kevin's workflow, revised 2026-09-09):
  - 可插入图片/ is a TRUSTED POOL curated by the user: only human-reviewed real
    assets (product screenshots, authorized case photos, company materials).
  - File names are semantic annotations, e.g. "GienWork-云端我的数字员工仪表盘.png".
    A "P{page}-" prefix is OPTIONAL; when present it serves only as a
    consistency hint (mismatch => warning).
  - The AI (not the user) scans the pool while designing pages, picks images by
    name semantics, reserves layout space in the SVG, and WRITES
    embed_images.json itself. Users never edit any manifest.

The pipeline's "zero bitmap information" iron rule is unaffected: body text and
diagrams stay native shapes; ONLY whitelisted evidence images become <p:pic>.

JSON schema (embed_images.json — authored by the AI, never by the user):
{
  "images_dir": "/abs/path/可插入图片",      # absolute, or relative to project_dir
  "images": [
    {"page": 12, "file": "GienWork-云端我的数字员工仪表盘.png",
     "x": 90, "y": 265, "w": 518, "h": 269,   # SVG px in 1280x720 canvas
     "border": "1B4593"}                      # optional hairline frame color
  ]
}
"""

from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

PX = 9525


def _inject(slide, img_path: Path, spec: dict) -> str:
    x, y = int(spec["x"] * PX), int(spec["y"] * PX)
    w, h = int(spec["w"] * PX), int(spec["h"] * PX)
    slide.shapes.add_picture(str(img_path), x, y, w, h)
    note = f"p{spec['page']} {img_path.name} {spec['w']}x{spec['h']}px"
    border = spec.get("border")
    if border:
        frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        frame.fill.background()
        frame.line.color.rgb = RGBColor.from_string(border.lstrip("#").upper())
        frame.line.width = Pt(1)
        frame.shadow.inherit = False
        note += " +frame"
    return note


def main() -> None:
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    pptx_path = Path(sys.argv[1])
    project_dir = Path(sys.argv[2])
    strict = "--strict" in sys.argv

    cfg_path = project_dir / "embed_images.json"
    if not cfg_path.is_file():
        print(f"[embed_images_inject] no embed_images.json in {project_dir} — nothing to inject")
        return
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    images = cfg.get("images", [])
    if not images:
        print("[embed_images_inject] empty images[] — nothing to inject")
        return

    images_dir = Path(cfg.get("images_dir", "可插入图片"))
    if not images_dir.is_absolute():
        images_dir = project_dir / images_dir
    if not images_dir.is_dir():
        msg = f"controlled images dir not found: {images_dir}"
        if strict:
            raise SystemExit(f"[embed_images_inject] {msg}")
        print(f"[WARN][embed_images_inject] {msg} — all entries skipped")
        return

    prs = Presentation(str(pptx_path))
    done, skipped = [], []
    for spec in images:
        page = int(spec["page"])
        name = os.path.basename(spec["file"])  # refuse any path component
        if name != spec["file"]:
            skipped.append(f"{spec['file']} (path components refused)")
            continue
        prefix = re.match(r"^P(\d+)-", name)
        if prefix and int(prefix.group(1)) != page:
            print(f"[WARN][embed_images_inject] {name}: filename page prefix P{prefix.group(1)}"
                  f" != declared page {page}")
        img_path = images_dir / name
        if not img_path.is_file():
            skipped.append(f"{name} (missing in {images_dir.name}/)")
            continue
        if not 1 <= page <= len(prs.slides):
            skipped.append(f"{name} (page {page} out of range)")
            continue
        done.append(_inject(prs.slides[page - 1], img_path, spec))

    prs.save(str(pptx_path))
    print(f"[embed_images_inject] injected {len(done)} evidence image(s):")
    for line in done:
        print(f"  + {line}")
    for line in skipped:
        print(f"  - SKIPPED {line}")
    if skipped and strict:
        raise SystemExit("[embed_images_inject] --strict: skipped entries present")
    print(f"[embed_images_inject] saved: {pptx_path}")


if __name__ == "__main__":
    main()
