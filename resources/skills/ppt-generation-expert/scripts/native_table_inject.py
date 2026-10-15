#!/usr/bin/env python3
"""native_table_inject.py — 注入原生可编辑表格（双击可改数）到已合并的册子里。

与 native_chart_inject.py 同一套契约：**SVG 里把表格区留白（什么都不画）**，表格数据与几何
写在 `<project_dir>/native_tables.json`，在 merge_into_template.py 之后、validate_pptx.py 之前运行。

为什么要有它：SVG 形状画的表格交付后**客户改不了数**（要改就得回到 SVG 重编译）。原生表格在
PowerPoint 里双击即可编辑单元格，行高列宽可拖，是"给客户/领导拿去改数"场景的正解。

JSON schema（native_tables.json）：
{
  "tables": [
    {
      "page": 13,                            # 1-based 页码
      "x": 70, "y": 454,                     # SVG px（1280x720 画布）
      "w": 1140, "h": 160,
      "col_widths": [340, 400, 400],          # px；总和须等于 w（不等则按比例归一 + 告警）
      "row_heights": 32,                      # px；数字=每行等高，数组=逐行指定
      "header": ["指标", "当前值", "建议目标"],  # 省略或 null = 无表头行
      "rows": [
        ["考核口径月合计", "8 月 90.60 万元", "9-12 月逐月递减至 65.37"]
      ],
      "header_fill": "0C2864", "header_color": "FFFFFF", "header_size": 14, "header_bold": true,
      "body_color": "0C2864", "body_size": 14, "body_fill": "FFFFFF",
      "align": ["left", "left", "left"],      # 逐列对齐（left|center|right）；缺省 left
      "banded": false, "band_fill": "E4E6EA", # 斑马行（默认关）
      "rule_color": "E4E6EA", "rule_width": 0.75,   # 行间细线
      "bold_col0": false,                     # 首列加粗（指标列常用）
      "emphasize": {"col": 2, "color": "1B4593", "bold": false}   # 强调某列正文（仅数据行，表头不受影响）
    }
  ]
}

必读注意：
- **python-pptx 默认表格样式是蓝色主题**（`{5C22544A-…}`），若不处理会把色板外颜色带进交付。
  本脚本强制把 tableStyleId 改写成 "No Style, No Grid" 并关掉 firstRow / bandRow，然后**逐格显式**
  设填充与边框——配色完全由本脚本控制，不受内置样式影响。
- SVG px → pt 的换算是 **×0.75**（96dpi→72pt），与导出器一致：14px = 10.5pt。
- 单元格文字过多时 PowerPoint 会自动增高行（行高是下限而非固定值）；表格因此可能高于 `h`，
  排版时给足余量。
- 依赖：Skill 锁定依赖对应的 Python 环境（含 python-pptx）。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from lxml import etree as _etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 9525                       # 1 SVG px = 1/96 inch
FONT = "Microsoft YaHei"
NO_STYLE_NO_GRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_LN_ORDER = ("a:lnL", "a:lnR", "a:lnT", "a:lnB")


def _rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.lstrip("#").upper())


def _emu(px: float) -> Emu:
    return Emu(int(round(px * PX)))


def _set_table_style(tbl, style_id: str = NO_STYLE_NO_GRID) -> None:
    tblPr = tbl._tbl.tblPr
    el = tblPr.find(qn("a:tableStyleId"))
    if el is None:
        el = _etree.SubElement(tblPr, qn("a:tableStyleId"))
    el.text = style_id


def _cell_borders(cell, color: str, width_pt: float, edges=("B",)) -> None:
    """给单元格加细线。a:tcPr 内 lnL/lnR/lnT/lnB 必须排在填充属性之前（schema 顺序）。"""
    tcPr = cell._tc.get_or_add_tcPr()
    ln_tags = [qn(t) for t in _LN_ORDER]
    for e in edges:
        tag = qn(f"a:ln{e}")
        for old in tcPr.findall(tag):
            tcPr.remove(old)
        ln = _etree.Element(tag)
        ln.set("w", str(int(round(width_pt * 12700))))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        fill = _etree.SubElement(ln, qn("a:solidFill"))
        clr = _etree.SubElement(fill, qn("a:srgbClr"))
        clr.set("val", color.lstrip("#").upper())
        pos = 0
        for i, child in enumerate(tcPr):
            if child.tag in ln_tags:
                pos = i + 1
        tcPr.insert(pos, ln)


def _fill_cell(cell, hexcolor: str) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = _rgb(hexcolor)


def _put_text(cell, text: str, size_px: float, color: str, bold: bool,
              align: str, margins: float = 8.0) -> None:
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = _emu(margins)
    cell.margin_right = _emu(margins)
    cell.margin_top = _emu(2)
    cell.margin_bottom = _emu(2)
    tf = cell.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    para.alignment = _ALIGN.get(align, PP_ALIGN.LEFT)
    run = para.add_run()
    run.text = text
    font = run.font
    font.name = FONT
    font.size = Pt(size_px * 0.75)          # px → pt（与导出器一致）
    font.bold = bold
    font.color.rgb = _rgb(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):            # 中文字体须显式声明 ea
        el = rPr.find(qn(tag))
        if el is None:
            el = _etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def _normalize_widths(widths: list[float], total: float, page: int) -> list[float]:
    s = sum(widths)
    if s <= 0:
        return [total / len(widths)] * len(widths)
    if abs(s - total) > 1.0:
        print(f"[WARN] p{page}: col_widths 合计 {s:g} ≠ w {total:g}，已按比例归一")
        return [w * total / s for w in widths]
    return widths


def _inject_table(slide, spec: dict) -> str:
    page = int(spec["page"])
    header = spec.get("header") or []
    rows = spec.get("rows") or []
    if not rows:
        raise SystemExit(f"native_tables.json p{page}: rows 为空")
    n_cols = len(header) if header else len(rows[0])
    for i, r in enumerate(rows, 1):
        if len(r) != n_cols:
            raise SystemExit(f"native_tables.json p{page}: 第 {i} 行有 {len(r)} 格，表头有 {n_cols} 列")
    n_rows = len(rows) + (1 if header else 0)

    w, h = float(spec["w"]), float(spec["h"])
    frame = slide.shapes.add_table(n_rows, n_cols, _emu(spec["x"]), _emu(spec["y"]), _emu(w), _emu(h))
    tbl = frame.table

    _set_table_style(tbl)
    tbl.first_row = False          # 关掉内置样式的表头强调
    tbl.horz_banding = False       # 关掉内置斑马行（改由我们自己控制）

    cw = spec.get("col_widths") or [w / n_cols] * n_cols
    cw = _normalize_widths([float(v) for v in cw], w, page)
    for i, px in enumerate(cw):
        tbl.columns[i].width = _emu(px)

    rh = spec.get("row_heights", (h / n_rows))
    heights = [float(rh)] * n_rows if isinstance(rh, (int, float)) else [float(v) for v in rh]
    for i, px in enumerate(heights[:n_rows]):
        tbl.rows[i].height = _emu(px)

    align = spec.get("align") or []
    aligns = [align[i] if i < len(align) else "left" for i in range(n_cols)]
    emph = spec.get("emphasize") or {}
    rule_color = spec.get("rule_color", "E4E6EA")
    rule_w = float(spec.get("rule_width", 0.75))

    def style_row(r_idx: int, values, bg: str, fg: str, size: float, bold: bool,
                  bold_col0: bool = False, is_header: bool = False) -> None:
        for c_idx, value in enumerate(values):
            cell = tbl.cell(r_idx, c_idx)
            _fill_cell(cell, bg)
            color, is_bold = fg, bold
            if emph and not is_header and c_idx == int(emph.get("col", -1)):
                color = emph.get("color", fg)
                is_bold = is_bold or bool(emph.get("bold", False))
            if bold_col0 and c_idx == 0:
                is_bold = True
            _put_text(cell, str(value), size, color, is_bold, aligns[c_idx])
            if r_idx > 0:                     # 首行不加顶线，其余行加细线
                _cell_borders(cell, rule_color, rule_w, edges=("T",))

    r0 = 0
    if header:
        style_row(0, header, spec.get("header_fill", "0C2864"),
                  spec.get("header_color", "FFFFFF"), float(spec.get("header_size", 14)),
                  bool(spec.get("header_bold", True)), is_header=True)
        r0 = 1
    band = bool(spec.get("banded", False))
    for i, row in enumerate(rows):
        bg = spec.get("band_fill", "E4E6EA") if (band and i % 2 == 1) else spec.get("body_fill", "FFFFFF")
        style_row(r0 + i, row, bg, spec.get("body_color", "0C2864"),
                  float(spec.get("body_size", 14)), False, bool(spec.get("bold_col0", False)))

    return f"p{page} table {n_rows}x{n_cols} {w:g}x{h:g} @({spec['x']:g},{spec['y']:g})"


def main() -> None:
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) < 2:
        raise SystemExit(__doc__)
    pptx_path, project_dir = Path(args[0]), Path(args[1])
    strict = "--strict" in sys.argv

    cfg_path = project_dir / "native_tables.json"
    if not cfg_path.is_file():
        print(f"[native_table_inject] no native_tables.json in {project_dir} — nothing to inject")
        return
    tables = json.loads(cfg_path.read_text(encoding="utf-8")).get("tables", [])
    if not tables:
        print("[native_table_inject] empty tables[] — nothing to inject")
        return

    prs = Presentation(str(pptx_path))
    done = []
    for spec in tables:
        page = int(spec["page"])
        if not 1 <= page <= len(prs.slides):
            msg = f"native_tables.json: page {page} out of range (1..{len(prs.slides)})"
            if strict:
                raise SystemExit(msg)
            print(f"[WARN] {msg} — skipped")
            continue
        done.append(_inject_table(prs.slides[page - 1], spec))

    prs.save(str(pptx_path))
    print(f"[native_table_inject] injected {len(done)} native table(s):")
    for line in done:
        print(f"  - {line}")
    print(f"[native_table_inject] saved: {pptx_path}")


if __name__ == "__main__":
    main()
